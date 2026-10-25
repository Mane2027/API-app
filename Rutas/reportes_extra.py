import io
from datetime import date, datetime, time, timedelta
from typing import Literal, Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from dependencies import get_db
from models import InventoryMovement, Product, Provider, Purchase, Sale
from Rutas.auth import require_role

router = APIRouter(prefix="/reportes", tags=["Reportes"])

ROLES_REPORTES = ("admin", "inventario")
Formato = Literal["excel", "pdf"]


def _rango(inicio: date, fin: date):
    if inicio > fin:
        raise HTTPException(status_code=400, detail="La fecha de inicio no puede ser mayor que la final")
    return datetime.combine(inicio, time.min), datetime.combine(fin + timedelta(days=1), time.min)


def _fmt(dt):
    return dt.strftime("%Y-%m-%d %H:%M") if dt else ""


def _respuesta(formato, nombre, titulo, columnas, filas, resumen=None):
    """Genera el archivo en memoria y lo devuelve como descarga."""
    resumen = resumen or []
    buffer = io.BytesIO()

    if formato == "excel":
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        ws = wb.active
        ws.append([titulo])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([])
        ws.append(columnas)
        for c in ws[3]:
            c.font = Font(bold=True)
        for f in filas:
            ws.append(list(f))
        if resumen:
            ws.append([])
            for etiqueta, valor in resumen:
                ws.append([etiqueta, valor])
                ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        for col in ws.iter_cols(min_row=3):
            ancho = max(len(str(c.value if c.value is not None else "")) for c in col) + 2
            ws.column_dimensions[col[0].column_letter].width = min(ancho, 50)
        wb.save(buffer)
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        extension = "xlsx"
    else:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        estilos = getSampleStyleSheet()
        doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), title=titulo)
        elementos = [Paragraph(titulo, estilos["Title"]), Spacer(1, 12)]
        data = [list(columnas)] + [[("" if v is None else str(v)) for v in f] for f in filas]
        if len(data) == 1:
            data.append(["Sin datos para este reporte"] + [""] * (len(columnas) - 1))
        tabla = Table(data, repeatRows=1)
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#7C3AED")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F3FF")]),
        ]))
        elementos.append(tabla)
        if resumen:
            elementos.append(Spacer(1, 12))
            for etiqueta, valor in resumen:
                elementos.append(Paragraph(f"<b>{etiqueta}:</b> {valor}", estilos["Normal"]))
        doc.build(elementos)
        media = "application/pdf"
        extension = "pdf"

    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type=media,
        headers={"Content-Disposition": f'attachment; filename="{nombre}.{extension}"'},
    )


@router.get("/ventas/rango", dependencies=[Depends(require_role(*ROLES_REPORTES))])
def reporte_ventas_rango(
    fecha_inicio: date,
    fecha_fin: date,
    formato: Formato = "excel",
    db: Session = Depends(get_db),
):
    desde, hasta = _rango(fecha_inicio, fecha_fin)
    ventas = (
        db.query(Sale).filter(Sale.date >= desde, Sale.date < hasta).order_by(Sale.date).all()
    )
    filas = [(v.id, _fmt(v.date), v.user, round(v.total, 2)) for v in ventas]
    resumen = [
        ("Número de ventas", len(ventas)),
        ("Total vendido", round(sum(v.total for v in ventas), 2)),
    ]
    return _respuesta(
        formato, f"ventas_{fecha_inicio}_{fecha_fin}",
        f"Ventas del {fecha_inicio} al {fecha_fin}",
        ["ID", "Fecha", "Usuario", "Total"], filas, resumen,
    )


@router.get("/compras/proveedor", dependencies=[Depends(require_role(*ROLES_REPORTES))])
def reporte_compras_proveedor(
    fecha_inicio: date,
    fecha_fin: date,
    provider_id: Optional[int] = None,
    formato: Formato = "excel",
    db: Session = Depends(get_db),
):
    desde, hasta = _rango(fecha_inicio, fecha_fin)
    consulta = (
        db.query(Purchase, Provider.name)
        .outerjoin(Provider, Purchase.provider_id == Provider.id)
        .filter(Purchase.date >= desde, Purchase.date < hasta)
    )
    if provider_id:
        consulta = consulta.filter(Purchase.provider_id == provider_id)
    registros = consulta.order_by(Provider.name, Purchase.date).all()

    filas, subtotales = [], {}
    for compra, proveedor in registros:
        nombre = proveedor or f"Proveedor #{compra.provider_id}"
        filas.append((nombre, compra.id, _fmt(compra.date), compra.user, round(compra.total, 2)))
        subtotales[nombre] = subtotales.get(nombre, 0) + compra.total

    resumen = [(f"Subtotal {n}", round(t, 2)) for n, t in subtotales.items()]
    resumen.append(("Total general", round(sum(subtotales.values()), 2)))
    return _respuesta(
        formato, f"compras_proveedor_{fecha_inicio}_{fecha_fin}",
        f"Compras por proveedor del {fecha_inicio} al {fecha_fin}",
        ["Proveedor", "Compra", "Fecha", "Usuario", "Total"], filas, resumen,
    )


@router.get("/stock/bajo", dependencies=[Depends(require_role(*ROLES_REPORTES))])
def reporte_stock_bajo(formato: Formato = "excel", db: Session = Depends(get_db)):
    productos = (
        db.query(Product)
        .filter(Product.is_active.is_(True), Product.stock <= Product.min_stock)
        .order_by(Product.stock, Product.name)
        .all()
    )
    filas = [
        (p.sku, p.name, p.brand, p.stock, p.min_stock, "Agotado" if (p.stock or 0) <= 0 else "Bajo")
        for p in productos
    ]
    agotados = sum(1 for p in productos if (p.stock or 0) <= 0)
    resumen = [("Productos en alerta", len(productos)), ("Agotados", agotados)]
    return _respuesta(
        formato, "stock_bajo", "Productos con stock bajo o agotado",
        ["SKU", "Nombre", "Marca", "Stock", "Mínimo", "Estado"], filas, resumen,
    )


@router.get("/movimientos", dependencies=[Depends(require_role(*ROLES_REPORTES))])
def reporte_movimientos(
    fecha_inicio: date,
    fecha_fin: date,
    product_id: Optional[int] = None,
    formato: Formato = "excel",
    db: Session = Depends(get_db),
):
    desde, hasta = _rango(fecha_inicio, fecha_fin)
    consulta = db.query(InventoryMovement).filter(
        InventoryMovement.timestamp >= desde, InventoryMovement.timestamp < hasta
    )
    if product_id:
        consulta = consulta.filter(InventoryMovement.product_id == product_id)
    movimientos = consulta.order_by(InventoryMovement.timestamp).all()

    nombres = {p.id: p.name for p in db.query(Product).all()}
    filas = [
        (
            _fmt(m.timestamp), nombres.get(m.product_id, f"Producto #{m.product_id}"),
            m.movement_type, m.quantity, m.previous_stock, m.new_stock, m.performed_by,
        )
        for m in movimientos
    ]
    entradas = sum(m.quantity for m in movimientos if m.movement_type == "entrada")
    salidas = sum(m.quantity for m in movimientos if m.movement_type == "salida")
    resumen = [("Unidades que entraron", entradas), ("Unidades que salieron", salidas)]
    return _respuesta(
        formato, f"movimientos_{fecha_inicio}_{fecha_fin}",
        f"Movimientos de inventario del {fecha_inicio} al {fecha_fin}",
        ["Fecha", "Producto", "Tipo", "Cantidad", "Stock anterior", "Stock nuevo", "Usuario"],
        filas, resumen,
    )


@router.get("/vencimientos", dependencies=[Depends(require_role(*ROLES_REPORTES))])
def reporte_vencimientos(
    dias: int = Query(30, ge=1, le=730),
    formato: Formato = "excel",
    db: Session = Depends(get_db),
):
    hoy = date.today()
    limite = hoy + timedelta(days=dias)
    productos = (
        db.query(Product)
        .filter(Product.expiration_date.isnot(None), Product.expiration_date <= limite)
        .order_by(Product.expiration_date)
        .all()
    )
    filas = []
    for p in productos:
        restantes = (p.expiration_date - hoy).days
        estado = "Vencido" if restantes < 0 else "Por vencer"
        filas.append((p.sku, p.name, p.brand, p.stock, str(p.expiration_date), restantes, estado))
    vencidos = sum(1 for f in filas if f[6] == "Vencido")
    resumen = [("Productos listados", len(filas)), ("Ya vencidos", vencidos)]
    return _respuesta(
        formato, f"vencimientos_{dias}_dias", f"Productos que vencen en los próximos {dias} días",
        ["SKU", "Nombre", "Marca", "Stock", "Vence", "Días restantes", "Estado"], filas, resumen,
    )
